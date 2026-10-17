import argparse
import csv
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = BASE_DIR / "data" / "sales_january.csv"
DEFAULT_OUTPUT = BASE_DIR / "output" / "sales_report.xlsx"

REQUIRED_COLUMNS = {
    "date", "order_id", "product", "quantity", "unit_price"
}


def load_sales(csv_path):
    records = []

    with csv_path.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)

        if not reader.fieldnames:
            raise ValueError("CSV file is empty or missing headers.")

        headers = [h.strip() for h in reader.fieldnames if h]
        if len(headers) != len(reader.fieldnames):
            raise ValueError("CSV contains an empty column header.")

        if len(headers) != len(set(headers)):
            raise ValueError("CSV contains duplicate column headers.")

        missing = REQUIRED_COLUMNS - set(headers)
        if missing:
            raise ValueError(
                "Missing required columns: " + ", ".join(sorted(missing))
            )

        for line_number, row in enumerate(reader, start=2):
            cleaned = {
                key.strip(): (value or "").strip()
                for key, value in row.items()
                if key is not None
            }

            if not any(cleaned.values()):
                continue

            try:
                sale_date = datetime.strptime(cleaned["date"], "%Y-%m-%d")
                quantity = int(cleaned["quantity"])
                unit_price = Decimal(cleaned["unit_price"])
            except (ValueError, InvalidOperation, KeyError):
                raise ValueError(
                    f"Invalid date, quantity, or price on CSV line {line_number}."
                )

            if quantity <= 0:
                raise ValueError(
                    f"Quantity must be positive on CSV line {line_number}."
                )

            if not unit_price.is_finite() or unit_price < 0:
                raise ValueError(
                    f"Unit price must be a valid non-negative number "
                    f"on CSV line {line_number}."
                )

            if not cleaned["order_id"] or not cleaned["product"]:
                raise ValueError(
                    f"Order ID and product are required on line {line_number}."
                )

            records.append({
                "date": sale_date,
                "order_id": cleaned["order_id"],
                "product": cleaned["product"],
                "quantity": quantity,
                "unit_price": unit_price,
                "total": unit_price * quantity,
            })

    if not records:
        raise ValueError("No valid sales records found.")

    return records


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="17365D")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        letter = get_column_letter(column_cells[0].column)
        ws.column_dimensions[letter].width = min(max_length + 3, 30)


def create_report(records, output_path):
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    details = workbook.create_sheet("Sales Details")
    summary = workbook.create_sheet("Product Summary")

    details.append([
        "Date", "Order ID", "Product", "Quantity",
        "Unit Price", "Total Sales"
    ])

    product_totals = defaultdict(
        lambda: {"quantity": 0, "sales": Decimal("0")}
    )
    total_units = 0
    total_sales = Decimal("0")

    for record in sorted(records, key=lambda item: item["date"]):
        details.append([
            record["date"],
            record["order_id"],
            record["product"],
            record["quantity"],
            float(record["unit_price"]),
            float(record["total"]),
        ])

        product = product_totals[record["product"]]
        product["quantity"] += record["quantity"]
        product["sales"] += record["total"]
        total_units += record["quantity"]
        total_sales += record["total"]

    for row in details.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        row[4].number_format = '#,##0.00'
        row[5].number_format = '#,##0.00'

    summary.append(["Product", "Units Sold", "Total Sales"])
    for product_name, values in sorted(product_totals.items()):
        summary.append([
            product_name,
            values["quantity"],
            float(values["sales"]),
        ])

    for row in summary.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = '#,##0.00'

    overview.append(["Sales Report", "Value"])
    overview.append(["Total Orders", len(records)])
    overview.append(["Unique Products", len(product_totals)])
    overview.append(["Units Sold", total_units])
    overview.append(["Total Sales", float(total_sales)])
    overview["B5"].number_format = '#,##0.00'

    for sheet in workbook.worksheets:
        style_sheet(sheet)

    overview.column_dimensions["A"].width = 22
    overview.column_dimensions["B"].width = 18

    workbook.save(output_path)
    return total_sales, total_units, len(records)


def main():
    parser = argparse.ArgumentParser(
        description="Generate an Excel sales report from a CSV file."
    )
    parser.add_argument(
        "--input", type=Path, default=DEFAULT_INPUT,
        help="Path to the input CSV file."
    )
    parser.add_argument(
        "--output", type=Path, default=DEFAULT_OUTPUT,
        help="Path for the generated Excel report."
    )
    args = parser.parse_args()

    try:
        records = load_sales(args.input)
        total_sales, units, orders = create_report(records, args.output)
    except (OSError, ValueError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1

    print("Sales report created successfully!")
    print(f"Orders: {orders}")
    print(f"Units sold: {units}")
    print(f"Total sales: {total_sales:.2f}")
    print(f"Output: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
